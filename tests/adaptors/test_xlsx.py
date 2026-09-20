import io
import zipfile
from datetime import datetime, time

import openpyxl
import pytest

from usali.adaptors.pdf import cluster_rows
from usali.adaptors.xlsx import XLSX_COL_STEP, XLSX_ROW_STEP, extract_words_from_xlsx_bytes


def _workbook(cells: dict[str, object]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Report"
    for ref, value in cells.items():
        ws[ref] = value
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_one_word_per_non_empty_cell_on_the_grid():
    words = extract_words_from_xlsx_bytes(_workbook({"A1": "Lakeside Test Lodge", "C2": 5, "B4": "x"}))
    by_text = {w.text: (w.x0, w.top) for w in words}
    assert by_text == {
        "Lakeside Test Lodge": (1 * XLSX_COL_STEP, 1 * XLSX_ROW_STEP),
        "5": (3 * XLSX_COL_STEP, 2 * XLSX_ROW_STEP),
        "x": (2 * XLSX_COL_STEP, 4 * XLSX_ROW_STEP),
    }


def test_cluster_rows_recovers_the_sheet_rows():
    words = extract_words_from_xlsx_bytes(_workbook({"A1": "a", "B1": "b", "A2": "c"}))
    rows = [[w.text for w in sorted(r, key=lambda w: w.x0)] for r in cluster_rows(words)]
    assert rows == [["a", "b"], ["c"]]


def test_numbers_dates_and_times_render_as_exact_text():
    words = extract_words_from_xlsx_bytes(
        _workbook({
            "A1": 361.63, "B1": 40699.1, "C1": 0, "D1": -5992.87, "E1": 700.0,
            "A2": datetime(2026, 8, 13, 0, 0), "B2": datetime(2026, 8, 13, 15, 12, 34),
            "C2": time(15, 12, 34), "D2": True,
        })
    )
    assert [w.text for w in sorted(words, key=lambda w: (w.top, w.x0))] == [
        "361.63", "40699.1", "0", "-5992.87", "700",
        "2026-08-13", "2026-08-13T15:12:34", "15:12:34", "TRUE",
    ]


def test_only_the_first_worksheet_is_read():
    wb = openpyxl.Workbook()
    first = wb.active
    assert first is not None
    first["A1"] = "first"
    wb.create_sheet("Second")["A1"] = "second"
    buf = io.BytesIO()
    wb.save(buf)
    assert [w.text for w in extract_words_from_xlsx_bytes(buf.getvalue())] == ["first"]


def test_not_a_workbook_raises():
    with pytest.raises(ValueError, match="not an XLSX workbook"):
        extract_words_from_xlsx_bytes(b"PK\x03\x04 not really a zip")


def test_cell_text_renders_a_whole_number_float_without_a_point():
    from usali.adaptors.xlsx import _cell_text

    assert _cell_text(700.0) == "700"
    assert _cell_text(150.5) == "150.5"
    assert _cell_text(361.63) == "361.63"


def test_a_workbook_with_corrupt_sheet_xml_is_refused_loudly():
    good = _workbook({"A1": "hello"})
    in_zip = zipfile.ZipFile(io.BytesIO(good))
    out_buf = io.BytesIO()
    with zipfile.ZipFile(out_buf, "w") as out_zip:
        for item in in_zip.infolist():
            content = in_zip.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                content = b"<not valid xml"
            out_zip.writestr(item, content)
    with pytest.raises(ValueError, match="not an XLSX workbook"):
        extract_words_from_xlsx_bytes(out_buf.getvalue())


def test_an_elapsed_time_cell_is_refused_not_guessed():
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Report"
    ws["A1"] = 1.5
    ws["A1"].number_format = "[h]:mm:ss"
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(ValueError, match="unsupported cell value type timedelta"):
        extract_words_from_xlsx_bytes(buf.getvalue())


def test_an_empty_workbook_yields_no_words():
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Report"
    buf = io.BytesIO()
    wb.save(buf)
    assert extract_words_from_xlsx_bytes(buf.getvalue()) == []
